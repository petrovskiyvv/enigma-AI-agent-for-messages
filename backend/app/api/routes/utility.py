import csv
import io

from fastapi import APIRouter
from fastapi.responses import StreamingResponse

from app.services.ticket_service import TicketService

router = APIRouter(tags=["utility"])
service = TicketService()

_CSV_FIELDS = [
    "id", "created_at", "full_name", "facility", "phone", "email",
    "device_numbers", "device_type", "emotional_tone", "category",
    "issue_summary", "status",
]

_XLSX_HEADERS = {
    "id": "ID",
    "created_at": "Дата",
    "full_name": "ФИО",
    "facility": "Объект",
    "phone": "Телефон",
    "email": "Email",
    "device_numbers": "Приборы",
    "device_type": "Тип прибора",
    "emotional_tone": "Тональность",
    "category": "Категория",
    "issue_summary": "Суть",
    "status": "Статус",
}


@router.get("/stats")
async def get_stats():
    return service.get_stats()


@router.get("/export/csv")
async def export_csv():
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=_CSV_FIELDS, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(service.get_list())
    output.seek(0)

    return StreamingResponse(
        io.BytesIO(output.getvalue().encode("utf-8-sig")),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=tickets_export.csv"},
    )


@router.get("/export/xlsx")
async def export_xlsx():
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        return StreamingResponse(
            io.BytesIO(b"openpyxl not installed"),
            status_code=500,
            media_type="text/plain",
        )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Обращения"

    # Шапка
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F3150", end_color="1F3150", fill_type="solid")

    for col_idx, field in enumerate(_CSV_FIELDS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=_XLSX_HEADERS.get(field, field))
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Данные
    for row_idx, ticket in enumerate(service.get_list(), start=2):
        for col_idx, field in enumerate(_CSV_FIELDS, start=1):
            ws.cell(row=row_idx, column=col_idx, value=ticket.get(field, ""))

    # Ширина колонок
    col_widths = [8, 18, 25, 22, 16, 24, 16, 20, 12, 16, 35, 12]
    for col_idx, width in enumerate(col_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=tickets_export.xlsx"},
    )